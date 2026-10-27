#!/usr/bin/env python3
"""Extract text, tables, images, and speaker notes from a .pptx file.

Usage:
    python3 extract-pptx.py <input.pptx> [output_dir]

Requires python-pptx >= 1.0:
    uv run --with python-pptx python extract-pptx.py deck.pptx deck/
    # or a prepared virtual environment

The output directory receives:
    extracted-slides.json
    assets/<image files>

Images are deduplicated by content hash. The script does not render the
presentation; SmartArt, charts, animations, and embedded video may be reduced
to text or omitted.
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
from typing import Any, Iterator

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:  # pragma: no cover - dependency error path
    sys.stderr.write(
        "error: python-pptx is required.\n"
        "Run with: uv run --with python-pptx python extract-pptx.py "
        "<input.pptx> [output_dir]\n"
    )
    raise SystemExit(2)


MAX_GROUP_DEPTH = 4


def iter_shapes(shapes, depth: int = 0) -> Iterator[Any]:
    """Yield shapes in document order, descending into groups."""
    for shape in shapes:
        is_group = (
            getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP
            and depth < MAX_GROUP_DEPTH
        )
        if is_group:
            yield from iter_shapes(shape.shapes, depth + 1)
        else:
            yield shape


def shape_alt_text(shape: Any) -> str:
    """Return the shape's alt text when the XML provides it."""
    try:
        return shape._element._nvXxPr.cNvPr.get("descr", "") or ""
    except Exception:
        return ""


def extract_pptx(file_path: str, output_dir: str) -> list[dict[str, Any]]:
    try:
        prs = Presentation(file_path)
    except Exception as exc:
        raise SystemExit(f"error: cannot open {file_path}: {exc}") from exc
    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)

    seen_images: dict[str, str] = {}
    slides_data: list[dict[str, Any]] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_data: dict[str, Any] = {
            "number": slide_num,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
            "unsupported": [],
        }

        title_shape = slide.shapes.title
        title_id = title_shape.shape_id if title_shape is not None else None

        for shape in iter_shapes(slide.shapes):
            # Pictures -----------------------------------------------------
            try:
                image = shape.image
            except Exception:
                image = None

            if image is not None:
                blob = image.blob
                digest = hashlib.sha256(blob).hexdigest()
                image_name = seen_images.get(digest)
                if image_name is None:
                    extension = image.ext or "png"
                    image_name = f"slide{slide_num}_img{len(seen_images) + 1}.{extension}"
                    with open(os.path.join(assets_dir, image_name), "wb") as handle:
                        handle.write(blob)
                    seen_images[digest] = image_name

                slide_data["images"].append(
                    {
                        "path": f"assets/{image_name}",
                        "width": int(shape.width or 0),
                        "height": int(shape.height or 0),
                        "left": int(shape.left or 0),
                        "top": int(shape.top or 0),
                        "alt": shape_alt_text(shape),
                    }
                )

            # Tables -------------------------------------------------------
            if getattr(shape, "has_table", False):
                table_rows = [
                    [cell.text or "" for cell in row.cells] for row in shape.table.rows
                ]
                slide_data["content"].append({"type": "table", "rows": table_rows})

            # Charts are reported, not reproduced --------------------------
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                try:
                    chart_title = chart.chart_title.text_frame.text
                except Exception:
                    chart_title = ""
                slide_data["content"].append(
                    {
                        "type": "chart",
                        "chart_type": str(chart.chart_type),
                        "title": chart_title,
                    }
                )
                slide_data["unsupported"].append(
                    f"chart on slide {slide_num} ({chart.chart_type})"
                )

            # Text ---------------------------------------------------------
            if getattr(shape, "has_text_frame", False):
                text = shape.text or ""
                if not text.strip():
                    continue
                if title_id is not None and shape.shape_id == title_id:
                    slide_data["title"] = text
                else:
                    slide_data["content"].append({"type": "text", "content": text})

        # Speaker notes ----------------------------------------------------
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                slide_data["notes"] = notes_frame.text or ""

        slides_data.append(slide_data)

    return slides_data


def main() -> int:
    if len(sys.argv) < 2:
        sys.stderr.write("Usage: python3 extract-pptx.py <input.pptx> [output_dir]\n")
        return 1

    input_file = os.path.abspath(sys.argv[1])
    output_dir = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.getcwd()

    if not os.path.isfile(input_file):
        sys.stderr.write(f"error: cannot read {input_file}\n")
        return 1
    if not input_file.lower().endswith(".pptx"):
        sys.stderr.write("error: expected a .pptx file (legacy .ppt is not supported)\n")
        return 1

    os.makedirs(output_dir, exist_ok=True)
    slides = extract_pptx(input_file, output_dir)

    output_path = os.path.join(output_dir, "extracted-slides.json")
    with open(output_path, "w", encoding="utf-8") as handle:
        json.dump(slides, handle, ensure_ascii=False, indent=2)

    print(f"Extracted {len(slides)} slides to {output_path}")
    for slide in slides:
        images = len(slide["images"])
        tables = sum(1 for item in slide["content"] if item["type"] == "table")
        notes = "notes" if slide["notes"] else "no notes"
        print(
            f"  Slide {slide['number']}: {slide['title'] or '(no title)'} — "
            f"{images} image(s), {tables} table(s), {notes}"
        )
        for warning in slide["unsupported"]:
            print(f"    warning: {warning} was not rendered")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
